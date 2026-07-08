from pathlib import Path
from typing import Any

import polars as pl
import yaml
from openpyxl import load_workbook

from app.desktop.models.data_source import DataSource
from app.desktop.models.desktop_pipeline_state import DesktopPipelineState
from app.desktop.models.workspace_state import WorkspaceState


class PipelineDesktopService:
    """Application service used by the desktop UI to manage queries and workspace sources."""

    def __init__(self) -> None:
        self.workspace = WorkspaceState()
        self.state = self.workspace.active_query

    def create_pipeline(self, name: str = "Consulta sin nombre") -> DesktopPipelineState:
        self.workspace.active_query = DesktopPipelineState(pipeline_name=name)
        self.state = self.workspace.active_query
        self.workspace.is_dirty = True
        return self.state

    def get_excel_sheet_names(self, path: Path) -> list[str]:
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Por ahora solo soportamos archivos Excel .xlsx.")

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = list(workbook.sheetnames)
        workbook.close()

        return sheets

    def _read_excel_safe(self, path: Path, sheet_name: str | None = None) -> pl.DataFrame:
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Por ahora solo soportamos archivos Excel .xlsx.")

        workbook = load_workbook(path, read_only=True, data_only=True)

        if sheet_name:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active

        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()

        if not rows:
            return pl.DataFrame()

        raw_headers = rows[0]
        headers: list[str] = []
        used_headers: set[str] = set()

        for index, value in enumerate(raw_headers):
            header = str(value).strip() if value is not None else f"Column_{index + 1}"

            if not header:
                header = f"Column_{index + 1}"

            original_header = header
            counter = 2

            while header in used_headers:
                header = f"{original_header}_{counter}"
                counter += 1

            used_headers.add(header)
            headers.append(header)

        records: list[dict[str, object]] = []

        for row in rows[1:]:
            record: dict[str, object] = {}

            for index, header in enumerate(headers):
                record[header] = row[index] if index < len(row) else None

            records.append(record)

        if not records:
            return pl.DataFrame({header: [] for header in headers})

        return pl.DataFrame(records, infer_schema_length=None, strict=False)

    def add_data_source(
        self,
        path: Path,
        sheet_name: str | None = None,
    ) -> DataSource:
        suffix = path.suffix.lower()

        if suffix == ".csv":
            source_type = "CSV"
        elif suffix == ".xlsx":
            source_type = "Excel"
        elif suffix == ".parquet":
            source_type = "Parquet"
        else:
            raise ValueError(f"Tipo de fuente no soportado: {suffix}")

        source_name = path.stem

        if sheet_name:
            source_name = f"{path.stem}_{sheet_name}"

        source = DataSource(
            name=source_name,
            path=path,
            source_type=source_type,
            sheet_name=sheet_name,
        )

        self.workspace.add_source(source)
        return source

    def activate_source_as_query(self, source_name: str) -> pl.DataFrame:
        source = self.workspace.get_source_by_name(source_name)

        if source is None:
            raise ValueError(f"No existe la fuente: {source_name}")

        self.state.source_path = source.path
        self.state.pipeline_name = source.name

        if source.source_type == "CSV":
            self.state.steps = [
                {
                    "type": "read_csv",
                    "name": "Leer CSV",
                    "config": {"path": str(source.path)},
                }
            ]
            dataframe = pl.read_csv(source.path)

        elif source.source_type == "Excel":
            self.state.steps = [
                {
                    "type": "read_excel",
                    "name": "Leer Excel",
                    "config": {
                        "path": str(source.path),
                        "sheet_name": source.sheet_name,
                    },
                }
            ]
            dataframe = self._read_excel_safe(source.path, source.sheet_name)

        elif source.source_type == "Parquet":
            self.state.steps = [
                {
                    "type": "read_parquet",
                    "name": "Leer Parquet",
                    "config": {"path": str(source.path)},
                }
            ]
            dataframe = pl.read_parquet(source.path)

        else:
            raise ValueError(f"Fuente no soportada: {source.source_type}")

        self.state.is_dirty = True
        self.workspace.is_dirty = True
        return dataframe

    def open_csv(self, path: Path) -> pl.DataFrame:
        source = self.add_data_source(path)
        return self.activate_source_as_query(source.name)

    def open_excel(self, path: Path) -> pl.DataFrame:
        sheets = self.get_excel_sheet_names(path)
        sheet_name = sheets[0] if sheets else None
        source = self.add_data_source(path, sheet_name)
        return self.activate_source_as_query(source.name)

    def add_export_step(self, output_path: Path) -> None:
        self.state.output_path = output_path

        export_step = {
            "type": "export_parquet",
            "name": "Exportar Parquet",
            "config": {"path": str(output_path)},
        }

        self.state.steps = [
            step for step in self.state.steps if step.get("type") != "export_parquet"
        ]

        self.state.steps.append(export_step)
        self.state.is_dirty = True
        self.workspace.is_dirty = True

    def build_pipeline_document(self) -> dict[str, Any]:
        return {
            "name": self.state.pipeline_name,
            "steps": self.state.steps,
        }

    def save_pipeline(self, path: Path) -> None:
        document = self.build_pipeline_document()
        path.parent.mkdir(parents=True, exist_ok=True)

        with path.open("w", encoding="utf-8") as file:
            yaml.safe_dump(document, file, sort_keys=False, allow_unicode=True)

        self.state.pipeline_path = path
        self.state.is_dirty = False
        self.workspace.is_dirty = False

    def load_pipeline(self, path: Path) -> DesktopPipelineState:
        with path.open("r", encoding="utf-8") as file:
            document = yaml.safe_load(file) or {}

        self.workspace.active_query = DesktopPipelineState(
            pipeline_path=path,
            pipeline_name=document.get("name", path.stem),
            steps=document.get("steps", []),
            is_dirty=False,
        )

        self.state = self.workspace.active_query
        return self.state

    def preview_current_source(self, limit: int = 500) -> pl.DataFrame:
        if self.state.source_path is None:
            return pl.DataFrame()

        first_step = self.state.steps[0] if self.state.steps else {}
        config = first_step.get("config", {})
        suffix = self.state.source_path.suffix.lower()

        if suffix == ".csv":
            return pl.read_csv(self.state.source_path).head(limit)

        if suffix == ".xlsx":
            return self._read_excel_safe(
                self.state.source_path,
                config.get("sheet_name"),
            ).head(limit)

        if suffix == ".parquet":
            return pl.read_parquet(self.state.source_path).head(limit)

        return pl.DataFrame()

    def add_transformation_step(self, step: dict[str, Any]) -> None:
        export_steps = [
            current_step
            for current_step in self.state.steps
            if current_step.get("type") == "export_parquet"
        ]

        self.state.steps = [
            current_step
            for current_step in self.state.steps
            if current_step.get("type") != "export_parquet"
        ]

        self.state.steps.append(step)
        self.state.steps.extend(export_steps)
        self.state.is_dirty = True
        self.workspace.is_dirty = True

    def apply_preview_transformations(self, dataframe: pl.DataFrame) -> pl.DataFrame:
        result = dataframe

        for step in self.state.steps:
            step_type = step.get("type")
            config = step.get("config", {})

            if step_type in {"read_csv", "read_excel", "read_parquet", "export_parquet"}:
                continue

            if step_type == "rename_columns":
                result = result.rename(config.get("columns", {}))

            elif step_type == "filter_rows":
                expression = config.get("expression")

                if expression:
                    result = result.sql(f"SELECT * FROM self WHERE {expression}")

            elif step_type == "select_columns":
                columns = config.get("columns", [])

                if columns:
                    result = result.select(columns)

            elif step_type == "drop_columns":
                columns = config.get("columns", [])

                if columns:
                    result = result.drop(columns)

        return result

    def run_pipeline_preview_mode(self) -> pl.DataFrame:
        dataframe = self.preview_current_source()
        return self.apply_preview_transformations(dataframe)

    def run_pipeline_until_step(self, step_index: int) -> pl.DataFrame:
        dataframe = self.preview_current_source()
        steps_to_apply = self.state.steps[: step_index + 1]

        for step in steps_to_apply:
            step_type = step.get("type")
            config = step.get("config", {})

            if step_type in {"read_csv", "read_excel", "read_parquet", "export_parquet"}:
                continue

            if step_type == "rename_columns":
                dataframe = dataframe.rename(config.get("columns", {}))

            elif step_type == "filter_rows":
                expression = config.get("expression")

                if expression:
                    dataframe = dataframe.sql(f"SELECT * FROM self WHERE {expression}")

            elif step_type == "select_columns":
                columns = config.get("columns", [])

                if columns:
                    dataframe = dataframe.select(columns)

            elif step_type == "drop_columns":
                columns = config.get("columns", [])

                if columns:
                    dataframe = dataframe.drop(columns)

        return dataframe

    def generate_python_script(self) -> str:
        lines = ["import polars as pl", "", ""]
        dataframe_name = "df"

        for step in self.state.steps:
            step_type = step.get("type")
            config = step.get("config", {})

            if step_type == "read_csv":
                path = config.get("path", "")
                lines.append(f'{dataframe_name} = pl.read_csv(r"{path}")')

            elif step_type == "read_excel":
                path = config.get("path", "")
                sheet_name = config.get("sheet_name")

                if sheet_name:
                    lines.append(
                        f'{dataframe_name} = pl.read_excel(r"{path}", sheet_name="{sheet_name}")'
                    )
                else:
                    lines.append(f'{dataframe_name} = pl.read_excel(r"{path}")')

            elif step_type == "read_parquet":
                path = config.get("path", "")
                lines.append(f'{dataframe_name} = pl.read_parquet(r"{path}")')

            elif step_type == "rename_columns":
                columns = config.get("columns", {})
                lines.append(f"{dataframe_name} = {dataframe_name}.rename({columns!r})")

            elif step_type == "filter_rows":
                expression = config.get("expression", "")
                lines.append(
                    f'{dataframe_name} = {dataframe_name}.sql("SELECT * FROM self WHERE {expression}")'
                )

            elif step_type == "select_columns":
                columns = config.get("columns", [])
                lines.append(f"{dataframe_name} = {dataframe_name}.select({columns!r})")

            elif step_type == "drop_columns":
                columns = config.get("columns", [])
                lines.append(f"{dataframe_name} = {dataframe_name}.drop({columns!r})")

            elif step_type == "export_parquet":
                path = config.get("path", "")
                lines.append(f'{dataframe_name}.write_parquet(r"{path}")')

        return "\n".join(lines)
